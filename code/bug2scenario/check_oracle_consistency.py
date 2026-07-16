#!/usr/bin/env python3
"""Compare Dataset.xlsx labels with error.json oracle events."""

from __future__ import annotations

import csv
import json
from pathlib import Path, PurePosixPath
from typing import Any
import zipfile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PROCESSED = ROOT / "processed"


def load_dataset() -> dict[int, dict[str, Any]]:
    workbook = load_workbook(ROOT / "Dataset.xlsx", data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header = [cell for cell in rows[0] if cell is not None]
    by_index: dict[int, dict[str, Any]] = {}
    for raw_row in rows[1:]:
        if raw_row[0] is None:
            continue
        row = {
            str(header[index]): raw_row[index]
            for index in range(len(header))
            if index < len(raw_row)
        }
        try:
            by_index[int(row["Index"])] = row
        except (TypeError, ValueError, KeyError):
            continue
    return by_index


def available_case_zips() -> list[Path]:
    zips = []
    for path in ROOT.glob("*.zip"):
        stem = path.stem
        if stem.isdigit():
            zips.append(path)
    return sorted(zips, key=lambda item: int(item.stem))


def read_error_json_from_zip(zip_path: Path) -> dict[str, Any] | None:
    inner_name = f"{zip_path.stem}/{zip_path.stem}_error.json"
    with zipfile.ZipFile(zip_path) as archive:
        names = set(archive.namelist())
        if inner_name not in names:
            return None
        return json.loads(archive.read(inner_name))


def derive_group(collision: int | None, stuck: int | None, lane: int | None, red: int | None) -> str:
    if None in {collision, stuck, lane, red}:
        return "unknown"
    parts = []
    if collision:
        parts.append("C")
    if stuck:
        parts.append("S")
    if lane:
        parts.append("L")
    if red:
        parts.append("R")
    return "+".join(parts) if parts else "none"


def to_int_label(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def status_and_actions(dataset_labels: dict[str, int | None], error_labels: dict[str, int | None]) -> tuple[str, str, str]:
    if all(value is None for value in error_labels.values()):
        return "missing_error_json", "error.json missing or unreadable", "locate_error_json_or_extract_again"

    mismatches = []
    dataset_only = []
    error_only = []
    matches = []

    for key in ["collision", "stuck", "lane_invasion", "red"]:
        dataset_value = dataset_labels.get(key)
        error_value = error_labels.get(key)
        if dataset_value is None or error_value is None:
            continue
        if dataset_value == error_value:
            matches.append(key)
        else:
            mismatches.append(f"{key}: dataset={dataset_value}, error_json={error_value}")
            if dataset_value == 1 and error_value == 0:
                dataset_only.append(key)
            elif dataset_value == 0 and error_value == 1:
                error_only.append(key)

    dataset_group = derive_group(
        dataset_labels.get("collision"),
        dataset_labels.get("stuck"),
        dataset_labels.get("lane_invasion"),
        dataset_labels.get("red"),
    )
    error_group = derive_group(
        error_labels.get("collision"),
        error_labels.get("stuck"),
        error_labels.get("lane_invasion"),
        error_labels.get("red"),
    )
    if dataset_group != "unknown" and error_group != "unknown" and dataset_group != error_group:
        mismatches.append(f"group: dataset={dataset_group}, error_json={error_group}")

    if not mismatches:
        return "consistent", "all comparable labels match", "no_action_needed"
    if dataset_only and not error_only:
        return "dataset_only", "; ".join(mismatches), "needs_manual_review"
    if error_only and not dataset_only:
        return "error_json_only", "; ".join(mismatches), "needs_manual_review"
    if matches and mismatches:
        return "conflict", "; ".join(mismatches), "needs_manual_review"
    if mismatches:
        return "conflict", "; ".join(mismatches), "needs_manual_review"
    return "unknown", "unable to compare labels", "inspect_case_metadata"


def build_rows() -> list[dict[str, Any]]:
    dataset = load_dataset()
    rows = []
    for zip_path in available_case_zips():
        case_id = int(zip_path.stem)
        dataset_row = dataset.get(case_id, {})
        error_json = read_error_json_from_zip(zip_path)
        events = (error_json or {}).get("events") or {}
        dataset_labels = {
            "collision": to_int_label(dataset_row.get("Collision")),
            "stuck": to_int_label(dataset_row.get("Stuck")),
            "lane_invasion": to_int_label(dataset_row.get("Lane Invasion")),
            "red": to_int_label(dataset_row.get("Red")),
        }
        error_labels = {
            "collision": int(bool(events.get("crash"))) if error_json is not None else None,
            "stuck": int(bool(events.get("stuck"))) if error_json is not None else None,
            "lane_invasion": int(bool(events.get("lane_invasion"))) if error_json is not None else None,
            "red": int(bool(events.get("red"))) if error_json is not None else None,
        }
        consistency_status, conflict_details, suggested_action = status_and_actions(dataset_labels, error_labels)
        row = {
            "case_id": f"case_{case_id:03d}",
            "dataset_collision": dataset_labels["collision"],
            "dataset_stuck": dataset_labels["stuck"],
            "dataset_lane_invasion": dataset_labels["lane_invasion"],
            "dataset_red": dataset_labels["red"],
            "dataset_group": dataset_row.get("Group", "unknown"),
            "error_json_collision": error_labels["collision"],
            "error_json_stuck": error_labels["stuck"],
            "error_json_lane_invasion": error_labels["lane_invasion"],
            "error_json_red": error_labels["red"],
            "error_json_group": derive_group(
                error_labels["collision"],
                error_labels["stuck"],
                error_labels["lane_invasion"],
                error_labels["red"],
            ),
            "consistency_status": consistency_status,
            "conflict_details": conflict_details,
            "suggested_action": suggested_action,
        }
        rows.append(row)
    return rows


def write_outputs(rows: list[dict[str, Any]]) -> None:
    PROCESSED.mkdir(parents=True, exist_ok=True)
    csv_path = PROCESSED / "oracle_consistency.csv"
    md_path = PROCESSED / "oracle_consistency.md"
    fieldnames = [
        "case_id",
        "dataset_collision",
        "dataset_stuck",
        "dataset_lane_invasion",
        "dataset_red",
        "dataset_group",
        "error_json_collision",
        "error_json_stuck",
        "error_json_lane_invasion",
        "error_json_red",
        "error_json_group",
        "consistency_status",
        "conflict_details",
        "suggested_action",
    ]
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    lines = [
        "# Oracle Consistency",
        "",
        "This sheet compares Dataset.xlsx labels against `error.json` oracle events. Conflicts are recorded only; labels are not overwritten automatically.",
        "",
        "| Case | Dataset Group | Error Group | Status | Conflict Details | Suggested Action |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for row in rows:
        lines.append(
            f"| {row['case_id']} | {row['dataset_group']} | {row['error_json_group']} | {row['consistency_status']} | {row['conflict_details']} | {row['suggested_action']} |"
        )
    lines.extend(
        [
            "",
            "## Status Meanings",
            "",
            "- `consistent`: Dataset.xlsx and error.json agree on all comparable fields.",
            "- `dataset_only`: Dataset.xlsx has positive labels not reflected in error.json.",
            "- `error_json_only`: error.json has positive labels not reflected in Dataset.xlsx.",
            "- `conflict`: mixed or partially contradictory evidence.",
            "- `missing_error_json`: error.json could not be found or parsed.",
            "- `unknown`: comparison was not possible.",
            "",
            "Cases marked `dataset_only`, `error_json_only`, or `conflict` should be treated as `needs_manual_review`.",
            "",
        ]
    )
    md_path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    write_outputs(build_rows())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
