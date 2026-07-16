#!/usr/bin/env python3
"""Build a reproducible preprocessing pipeline for DriveFuzz/Autoware failure cases."""

from __future__ import annotations

import csv
from dataclasses import dataclass
import json
import math
import os
from pathlib import Path, PurePosixPath
import shutil
import subprocess
import sys
from typing import Any
import zipfile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PROCESSED = ROOT / "processed"
EXTRACTED = PROCESSED / "extracted"
CONVERTED = PROCESSED / "converted"
CASE_CARDS = PROCESSED / "case_cards"
LLM_INPUTS = PROCESSED / "llm_inputs"
PROMPTS = PROCESSED / "prompts"

TARGET_CASES = [1, 7, 3]
INSPECT_ONLY_CASES = [13, 24, 27, 37]
ZIP_CASES = TARGET_CASES + INSPECT_ONLY_CASES

FAULT_GROUP_MEANINGS = {
    "C": "Collision",
    "R": "Red-light violation",
    "L": "Lane invasion",
    "S": "Stuck / immobility",
    "C+L": "Collision + lane invasion",
    "L+R": "Lane invasion + red-light violation",
    "C+R": "Collision + red-light violation",
    "C+L+R": "Collision + lane invasion + red-light violation",
}

TOPIC_SPECS = [
    {"topic_name": "/carla/ego_vehicle/imu/imu", "keywords": ["carla__ego_vehicle__imu__imu", "imu1"], "kind": "imu"},
    {"topic_name": "/image_raw", "keywords": ["image_raw"], "kind": "image"},
    {"topic_name": "/points_raw", "keywords": ["points_raw"], "kind": "pointcloud"},
    {"topic_name": "/detection/fusion_tools/objects", "keywords": ["detection__fusion_tools__objects", "fusion_tools", "objects"], "kind": "object_array"},
    {"topic_name": "/current_pose", "keywords": ["current_pose"], "kind": "pose"},
    {"topic_name": "/prediction/motion_predictor/objects", "keywords": ["prediction__motion_predictor__objects", "motion_predictor", "objects"], "kind": "object_array"},
    {"topic_name": "/lane_waypoints_array", "keywords": ["lane_waypoints_array"], "kind": "lane_array"},
    {"topic_name": "/final_waypoints", "keywords": ["final_waypoints"], "kind": "final_waypoints"},
    {"topic_name": "/vehicle_cmd", "keywords": ["vehicle_cmd"], "kind": "vehicle_cmd"},
    {"topic_name": "/carla/ego_vehicle/vehicle_status", "keywords": ["carla__ego_vehicle__vehicle_status", "vehicle_status"], "kind": "vehicle_status"},
    {"topic_name": "/carla/ego_vehicle/collision", "keywords": ["carla__ego_vehicle__collision", "collision"], "kind": "collision"},
    {"topic_name": "/carla/ego_vehicle/odometry", "keywords": ["carla__ego_vehicle__odometry", "odometry"], "kind": "odometry"},
    {"topic_name": "/carla/objects", "keywords": ["carla__objects", "objects"], "kind": "object_array"},
    {"topic_name": "/carla/traffic_lights", "keywords": ["carla__traffic_lights", "traffic_lights"], "kind": "traffic_lights"},
]

CSV_SAMPLE_LIMIT = 200


@dataclass
class TopicMatch:
    topic_name: str
    found: bool
    matched_file_path: str
    file_type: str
    notes: str
    kind: str
    actual_topic: str | None = None
    msg_type: str | None = None


def ensure_dirs() -> None:
    for path in [PROCESSED, EXTRACTED, CONVERTED, CASE_CARDS, LLM_INPUTS, PROMPTS]:
        path.mkdir(parents=True, exist_ok=True)


def human_size(num_bytes: int) -> str:
    units = ["B", "KB", "MB", "GB", "TB"]
    value = float(num_bytes)
    for unit in units:
        if value < 1024 or unit == units[-1]:
            return f"{value:.1f} {unit}" if unit != "B" else f"{int(value)} B"
        value /= 1024
    return f"{num_bytes} B"


def set_csv_field_limit() -> None:
    limit = sys.maxsize
    while True:
        try:
            csv.field_size_limit(limit)
            return
        except OverflowError:
            limit //= 10


def parse_json_cell(value: Any, default: Any = None) -> Any:
    if value in (None, ""):
        return default
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return default
        return json.loads(value)
    return value


def safe_float(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def timestamp_to_str(sec: Any, nsec: Any) -> str:
    if sec in (None, "") and nsec in (None, ""):
        return "unknown"
    return f"{sec}.{str(nsec).zfill(9)}"


def row_timestamp_str(row: dict[str, Any]) -> str:
    return timestamp_to_str(row.get("timestamp_sec"), row.get("timestamp_nsec"))


def normalize_topic(topic: str) -> str:
    return topic.strip("/").replace("/", "__")


def load_dataset() -> tuple[list[dict[str, Any]], dict[int, dict[str, Any]], str]:
    workbook = load_workbook(ROOT / "Dataset.xlsx", data_only=True)
    all_rows: list[dict[str, Any]] = []
    by_index: dict[int, dict[str, Any]] = {}
    md_lines = [
        "# Dataset Summary",
        "",
        f"- Workbook: `Dataset.xlsx`",
        f"- Sheets: {', '.join(workbook.sheetnames)}",
        "",
    ]

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header = list(rows[0])
        md_lines.append(f"## {sheet.title}")
        md_lines.append("")
        md_lines.append(f"- Rows including header: {len(rows)}")
        md_lines.append(f"- Columns: {', '.join(str(col) for col in header if col is not None)}")
        required = ["Index", "Collision", "Stuck", "Lane Invasion", "Red", "Sum", "Group", "Comment"]
        present = [col for col in required if col in header]
        missing = [col for col in required if col not in header]
        md_lines.append(f"- Required columns present: {', '.join(present)}")
        md_lines.append(f"- Missing required columns: {', '.join(missing) if missing else 'none'}")
        md_lines.append("")
        for row in rows[1:]:
            record = {str(header[i]): row[i] for i in range(len(header)) if header[i] is not None}
            if not any(value is not None for value in record.values()):
                continue
            record["sheet_name"] = sheet.title
            index_value = record.get("Index")
            zip_file = None
            if index_value is not None:
                try:
                    zip_file = f"{int(index_value)}.zip"
                    record["zip_file"] = zip_file
                except (TypeError, ValueError):
                    pass
            all_rows.append(record)
            if index_value is not None:
                try:
                    by_index[int(index_value)] = record
                except (TypeError, ValueError):
                    pass

    target_rows = []
    for case_id in ZIP_CASES:
        record = by_index.get(case_id)
        if not record:
            continue
        target_rows.append(
            "| {idx} | {zip_file} | {group} | {c} | {s} | {l} | {r} | {comment} |".format(
                idx=case_id,
                zip_file=record.get("zip_file", f"{case_id}.zip"),
                group=record.get("Group", "unknown"),
                c=int(record.get("Collision", 0) or 0),
                s=int(record.get("Stuck", 0) or 0),
                l=int(record.get("Lane Invasion", 0) or 0),
                r=int(record.get("Red", 0) or 0),
                comment=record.get("Comment", ""),
            )
        )
    md_lines.extend(
        [
            "## Target Cases",
            "",
            "| Index | Zip | Group | Collision | Stuck | Lane Invasion | Red | Comment |",
            "| --- | --- | --- | --- | --- | --- | --- | --- |",
            *target_rows,
            "",
            "## Group Meanings",
            "",
        ]
    )
    for key, meaning in FAULT_GROUP_MEANINGS.items():
        md_lines.append(f"- `{key}` = {meaning}")
    md_lines.append("")
    return all_rows, by_index, "\n".join(md_lines)


def write_dataset_summary(rows: list[dict[str, Any]], markdown: str) -> None:
    csv_path = PROCESSED / "dataset_summary.csv"
    md_path = PROCESSED / "dataset_summary.md"
    fieldnames = []
    seen = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.add(key)
                fieldnames.append(key)
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    md_path.write_text(markdown, encoding="utf-8")


def inspect_zip_file(zip_path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(zip_path) as archive:
        infos = [info for info in archive.infolist() if not info.is_dir()]
        top_levels: dict[str, int] = {}
        extensions: dict[str, int] = {}
        for info in infos:
            parts = [part for part in PurePosixPath(info.filename).parts if part]
            top = parts[0] if parts else "."
            top_levels[top] = top_levels.get(top, 0) + 1
            suffix = PurePosixPath(info.filename).suffix or "[no_ext]"
            extensions[suffix] = extensions.get(suffix, 0) + 1
        return {
            "file": zip_path.name,
            "size_bytes": zip_path.stat().st_size,
            "entries": len(infos),
            "top_levels": top_levels,
            "extensions": extensions,
            "samples": [info.filename for info in infos[:20]],
        }


def write_directory_and_zip_summary() -> None:
    relevant = sorted(
        [
            path
            for path in ROOT.iterdir()
            if path.is_file()
            and (path.suffix.lower() == ".zip" or path.name in {"Dataset.xlsx", "convert_sensor_data_field.py", "Topics selected.png"})
        ],
        key=lambda p: p.name,
    )
    lines = [
        "# Directory And Zip Summary",
        "",
        f"- Working directory: `{ROOT}`",
        "",
        "## Files",
        "",
        "| Name | Size |",
        "| --- | --- |",
    ]
    for path in relevant:
        lines.append(f"| {path.name} | {human_size(path.stat().st_size)} |")

    lines.extend(["", "## Zip Structures", ""])
    for case_id in ZIP_CASES:
        zip_path = ROOT / f"{case_id}.zip"
        info = inspect_zip_file(zip_path)
        lines.append(f"### {zip_path.name}")
        lines.append("")
        lines.append(f"- Size: {human_size(info['size_bytes'])}")
        lines.append(f"- File entries: {info['entries']}")
        lines.append(f"- Top-level folders: {json.dumps(info['top_levels'], ensure_ascii=False)}")
        lines.append(f"- File types: {json.dumps(info['extensions'], ensure_ascii=False)}")
        lines.append("- Sample entries:")
        for sample in info["samples"]:
            lines.append(f"  - `{sample}`")
        lines.append("")
    (PROCESSED / "directory_and_zip_summary.md").write_text("\n".join(lines), encoding="utf-8")


def extract_selected_cases() -> None:
    for case_id in TARGET_CASES:
        dest = EXTRACTED / f"case_{case_id:03d}"
        manifest_path = dest / "manifest.json"
        if manifest_path.exists():
            continue
        dest.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(ROOT / f"{case_id}.zip") as archive:
            for member in archive.infolist():
                if member.is_dir():
                    continue
                parts = [part for part in PurePosixPath(member.filename).parts if part]
                relative_parts = parts[1:] if len(parts) > 1 else parts
                relative = Path(*relative_parts)
                output_path = dest / relative
                output_path.parent.mkdir(parents=True, exist_ok=True)
                with archive.open(member) as source, output_path.open("wb") as target:
                    shutil.copyfileobj(source, target)


def load_case_static(case_id: int) -> dict[str, Any]:
    case_dir = EXTRACTED / f"case_{case_id:03d}"
    manifest = json.loads((case_dir / "manifest.json").read_text(encoding="utf-8"))
    connections = json.loads((case_dir / "connections.json").read_text(encoding="utf-8"))
    error = json.loads((case_dir / f"{case_id}_error.json").read_text(encoding="utf-8"))
    connection_map = {int(item["index"]): item for item in connections}
    return {
        "case_id": case_id,
        "case_dir": case_dir,
        "manifest": manifest,
        "connections": connections,
        "connection_map": connection_map,
        "error": error,
    }


def short_json(value: Any, max_len: int = 200) -> str:
    text = json.dumps(value, ensure_ascii=False)
    return text if len(text) <= max_len else text[: max_len - 3] + "..."


def csv_preview(path: Path) -> dict[str, Any]:
    set_csv_field_limit()
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader, [])
        first_row = next(reader, [])
    preview = []
    for cell in first_row[:8]:
        cell_text = str(cell)
        if len(cell_text) > 80:
            preview.append(cell_text[:77] + "...")
        else:
            preview.append(cell_text)
    return {"header": header, "sample": preview}


def write_file_inventory(cases: list[dict[str, Any]]) -> None:
    lines = [
        "# File Inventory",
        "",
        "This inventory only reads headers, small samples, and JSON keys. It does not load raw image arrays or full point clouds into memory.",
        "",
    ]
    for case in cases:
        case_id = case["case_id"]
        lines.append(f"## case_{case_id:03d}")
        lines.append("")
        lines.append(f"- Extracted path: `{case['case_dir']}`")
        lines.append(f"- Source zip: `{case_id}.zip`")
        lines.append(f"- Manifest messages: {case['manifest'].get('messages', 'unknown')}")
        lines.append("")
        lines.append("| Path | Type | Size | Topic / Keys | Header / Sample |")
        lines.append("| --- | --- | --- | --- | --- |")
        for path in sorted(case["case_dir"].rglob("*")):
            if path.is_dir():
                continue
            rel = path.relative_to(case["case_dir"]).as_posix()
            suffix = path.suffix.lower()
            size = human_size(path.stat().st_size)
            if suffix == ".json":
                data = json.loads(path.read_text(encoding="utf-8"))
                if isinstance(data, dict):
                    topic_or_keys = ", ".join(list(data.keys())[:12])
                elif isinstance(data, list):
                    topic_or_keys = f"list[{len(data)}]"
                else:
                    topic_or_keys = type(data).__name__
                sample = short_json(data, 180)
            elif suffix == ".csv":
                preview = csv_preview(path)
                topic_or_keys = "csv"
                if path.parent.name == "topics":
                    prefix = path.stem.split("_", 1)[0]
                    if prefix.isdigit() and int(prefix) in case["connection_map"]:
                        conn = case["connection_map"][int(prefix)]
                        topic_or_keys = f"{conn['topic']} ({conn['type']})"
                sample = f"header={preview['header'][:8]}; sample={preview['sample']}"
            else:
                topic_or_keys = suffix or "file"
                sample = ""
            lines.append(f"| {rel} | {suffix or 'file'} | {size} | {topic_or_keys} | {sample} |")
        lines.append("")
    (PROCESSED / "file_inventory.md").write_text("\n".join(lines), encoding="utf-8")


def find_topic_match(case: dict[str, Any], spec: dict[str, Any]) -> TopicMatch:
    topic_name = spec["topic_name"]
    normalized = normalize_topic(topic_name)
    topics_dir = case["case_dir"] / "topics"
    for connection in case["connections"]:
        actual_topic = connection["topic"]
        topic_ok = actual_topic == topic_name
        if topic_name == "/carla/ego_vehicle/imu/imu" and actual_topic.startswith("/carla/ego_vehicle/imu/imu"):
            topic_ok = True
        if topic_ok:
            csv_name = f"{int(connection['index']):04d}_{normalize_topic(actual_topic)}.csv"
            path = topics_dir / csv_name
            if not path.exists():
                candidates = sorted(topics_dir.glob(f"{int(connection['index']):04d}_*.csv"))
                path = candidates[0] if candidates else path
            return TopicMatch(
                topic_name=topic_name,
                found=path.exists(),
                matched_file_path=str(path.relative_to(case["case_dir"])) if path.exists() else "",
                file_type=path.suffix.lstrip(".") if path.exists() else "",
                notes="exact connection topic match" if actual_topic == topic_name else f"matched connection topic `{actual_topic}`",
                kind=spec["kind"],
                actual_topic=actual_topic,
                msg_type=connection.get("type"),
            )

    lowered_keywords = [keyword.lower() for keyword in spec["keywords"]]
    candidates = []
    for path in topics_dir.glob("*.csv"):
        rel = path.relative_to(case["case_dir"]).as_posix().lower()
        if any(keyword in rel for keyword in lowered_keywords):
            candidates.append(path)
    if candidates:
        path = sorted(candidates)[0]
        return TopicMatch(
            topic_name=topic_name,
            found=True,
            matched_file_path=str(path.relative_to(case["case_dir"])),
            file_type=path.suffix.lstrip("."),
            notes="fuzzy filename match",
            kind=spec["kind"],
        )

    return TopicMatch(
        topic_name=topic_name,
        found=False,
        matched_file_path="",
        file_type="",
        notes="not found via connection topic or fuzzy filename match",
        kind=spec["kind"],
    )


def write_topic_availability(cases: list[dict[str, Any]]) -> dict[int, dict[str, TopicMatch]]:
    output_path = PROCESSED / "topic_availability.csv"
    fieldnames = ["case_id", "topic_name", "found_or_not", "matched_file_path", "file_type", "notes"]
    matches_by_case: dict[int, dict[str, TopicMatch]] = {}
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for case in cases:
            case_matches: dict[str, TopicMatch] = {}
            for spec in TOPIC_SPECS:
                match = find_topic_match(case, spec)
                case_matches[spec["topic_name"]] = match
                writer.writerow(
                    {
                        "case_id": f"case_{case['case_id']:03d}",
                        "topic_name": spec["topic_name"],
                        "found_or_not": "found" if match.found else "not_found",
                        "matched_file_path": match.matched_file_path,
                        "file_type": match.file_type,
                        "notes": match.notes,
                    }
                )
            matches_by_case[case["case_id"]] = case_matches
    return matches_by_case


def read_last_nonempty_line(path: Path) -> str | None:
    with path.open("rb") as handle:
        handle.seek(0, os.SEEK_END)
        end = handle.tell()
        if end == 0:
            return None
        buffer = b""
        step = 8192
        position = end
        while position > 0:
            size = min(step, position)
            position -= size
            handle.seek(position)
            buffer = handle.read(size) + buffer
            lines = buffer.splitlines()
            if len(lines) >= 2 or position == 0:
                for line in reversed(lines):
                    if line.strip():
                        return line.decode("utf-8")
        return None


def parse_csv_line(header_line: str, data_line: str) -> dict[str, str]:
    reader = csv.DictReader([header_line, data_line])
    return next(reader)


def first_and_last_rows(path: Path) -> tuple[list[str], dict[str, str] | None, dict[str, str] | None]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        header_line = handle.readline()
        first_data_line = None
        for line in handle:
            if line.strip():
                first_data_line = line
                break
    if not header_line:
        return [], None, None
    header = next(csv.reader([header_line.rstrip("\n")]))
    if first_data_line is None:
        return header, None, None
    first_row = parse_csv_line(header_line, first_data_line)
    last_line = read_last_nonempty_line(path)
    if last_line is None or last_line.strip() == first_data_line.strip():
        return header, first_row, first_row
    last_row = parse_csv_line(header_line, last_line)
    return header, first_row, last_row


def scan_rows(path: Path, limit: int | None = None) -> tuple[list[str], list[dict[str, str]], bool]:
    set_csv_field_limit()
    rows: list[dict[str, str]] = []
    truncated = False
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = reader.fieldnames or []
        for index, row in enumerate(reader):
            if limit is not None and index >= limit:
                truncated = True
                break
            rows.append(row)
    return fieldnames, rows, truncated


def summarize_image_csv(path: Path, conversion_root: Path | None) -> dict[str, Any]:
    header, first_row, last_row = first_and_last_rows(path)
    if not first_row:
        return {"exists": True, "status": "empty_csv"}
    output_status = "not_attempted"
    decoded_files = []
    if conversion_root and conversion_root.exists():
        decoded_files = sorted(p.name for p in conversion_root.glob("*.npy"))
        if decoded_files:
            output_status = f"decoded {len(decoded_files)} frame(s)"
    data_obj = parse_json_cell(first_row.get("data"), {})
    return {
        "exists": True,
        "header": header,
        "encoding": first_row.get("encoding"),
        "height": int(first_row.get("height", 0) or 0),
        "width": int(first_row.get("width", 0) or 0),
        "step": int(first_row.get("step", 0) or 0),
        "first_timestamp": row_timestamp_str(first_row),
        "last_timestamp": row_timestamp_str(last_row or first_row),
        "bytes_per_frame_first_row": int(data_obj.get("bytes", 0) or 0),
        "decode_status": output_status,
        "decoded_sample_files": decoded_files[:5],
        "decode_error_preview": first_row.get("__decode_error") or "",
    }


def summarize_pointcloud_csv(path: Path, conversion_root: Path | None) -> dict[str, Any]:
    header, first_row, last_row = first_and_last_rows(path)
    if not first_row:
        return {"exists": True, "status": "empty_csv"}
    fields = parse_json_cell(first_row.get("fields"), []) or []
    output_status = "not_attempted"
    decoded_files = []
    if conversion_root and conversion_root.exists():
        decoded_files = sorted(p.name for p in conversion_root.glob("*.npy"))
        if decoded_files:
            output_status = f"decoded {len(decoded_files)} frame(s)"
    data_obj = parse_json_cell(first_row.get("data"), {})
    return {
        "exists": True,
        "header": header,
        "height": int(first_row.get("height", 0) or 0),
        "width": int(first_row.get("width", 0) or 0),
        "point_step": int(first_row.get("point_step", 0) or 0),
        "row_step": int(first_row.get("row_step", 0) or 0),
        "first_timestamp": row_timestamp_str(first_row),
        "last_timestamp": row_timestamp_str(last_row or first_row),
        "fields": [field.get("name") for field in fields if isinstance(field, dict)],
        "bytes_per_frame_first_row": int(data_obj.get("bytes", 0) or 0),
        "decode_status": output_status,
        "decoded_sample_files": decoded_files[:5],
        "decode_error_preview": first_row.get("__decode_error") or "",
    }


def summarize_imu_csv(path: Path) -> dict[str, Any]:
    _fieldnames, rows, truncated = scan_rows(path, limit=CSV_SAMPLE_LIMIT)
    if not rows:
        return {"exists": True, "rows_sampled": 0}
    angular_values = []
    linear_values = []
    for row in rows:
        angular = parse_json_cell(row.get("angular_velocity"), {}) or {}
        linear = parse_json_cell(row.get("linear_acceleration"), {}) or {}
        if isinstance(angular, dict):
            angular_values.append((
                safe_float(angular.get("x")),
                safe_float(angular.get("y")),
                safe_float(angular.get("z")),
            ))
        if isinstance(linear, dict):
            linear_values.append((
                safe_float(linear.get("x")),
                safe_float(linear.get("y")),
                safe_float(linear.get("z")),
            ))
    header, _first_row, last_row = first_and_last_rows(path)
    return {
        "exists": True,
        "header": header,
        "rows_sampled": len(rows),
        "sample_truncated": truncated,
        "first_timestamp": row_timestamp_str(rows[0]),
        "last_timestamp": row_timestamp_str(last_row or rows[-1]),
        "angular_velocity_z_range": min_max([triple[2] for triple in angular_values if triple[2] is not None]),
        "linear_acceleration_z_range": min_max([triple[2] for triple in linear_values if triple[2] is not None]),
    }


def min_max(values: list[float]) -> dict[str, float] | str:
    cleaned = [value for value in values if value is not None]
    if not cleaned:
        return "unknown"
    return {"min": min(cleaned), "max": max(cleaned)}


def summarize_pose_csv(path: Path) -> dict[str, Any]:
    _fieldnames, rows, truncated = scan_rows(path, limit=CSV_SAMPLE_LIMIT)
    if not rows:
        return {"exists": True, "rows_sampled": 0}
    positions = []
    for row in rows:
        pose = parse_json_cell(row.get("pose"), {}) or {}
        position = ((pose.get("position") or {}) if isinstance(pose, dict) else {})
        if isinstance(position, dict):
            positions.append(
                {
                    "x": safe_float(position.get("x")),
                    "y": safe_float(position.get("y")),
                    "z": safe_float(position.get("z")),
                }
            )
    _header, _first, last_row = first_and_last_rows(path)
    return {
        "exists": True,
        "rows_sampled": len(rows),
        "sample_truncated": truncated,
        "first_timestamp": row_timestamp_str(rows[0]),
        "last_timestamp": row_timestamp_str(last_row or rows[-1]),
        "start_position": positions[0] if positions else "unknown",
        "end_position_sample": positions[-1] if positions else "unknown",
    }


def summarize_vehicle_status_csv(path: Path) -> dict[str, Any]:
    _fieldnames, rows, truncated = scan_rows(path, limit=CSV_SAMPLE_LIMIT)
    velocities = []
    throttles = []
    brakes = []
    steers = []
    for row in rows:
        velocity = safe_float(row.get("velocity"))
        if velocity is not None:
            velocities.append(velocity)
        control = parse_json_cell(row.get("control"), {}) or {}
        if isinstance(control, dict):
            for source, target in [("throttle", throttles), ("brake", brakes), ("steer", steers)]:
                value = safe_float(control.get(source))
                if value is not None:
                    target.append(value)
    _header, _first, last_row = first_and_last_rows(path)
    return {
        "exists": True,
        "rows_sampled": len(rows),
        "sample_truncated": truncated,
        "first_timestamp": row_timestamp_str(rows[0]) if rows else "unknown",
        "last_timestamp": row_timestamp_str(last_row or rows[-1]) if rows else "unknown",
        "velocity_range": min_max(velocities),
        "throttle_range": min_max(throttles),
        "brake_range": min_max(brakes),
        "steer_range": min_max(steers),
    }


def summarize_vehicle_cmd_csv(path: Path) -> dict[str, Any]:
    _fieldnames, rows, truncated = scan_rows(path, limit=CSV_SAMPLE_LIMIT)
    steer_values = []
    accel_values = []
    brake_values = []
    target_speed_values = []
    for row in rows:
        steer_cmd = parse_json_cell(row.get("steer_cmd"), {}) or {}
        accel_cmd = parse_json_cell(row.get("accel_cmd"), {}) or {}
        brake_cmd = parse_json_cell(row.get("brake_cmd"), {}) or {}
        ctrl_cmd = parse_json_cell(row.get("ctrl_cmd"), {}) or {}
        for value, target in [
            (safe_float((steer_cmd or {}).get("steer")), steer_values),
            (safe_float((accel_cmd or {}).get("accel")), accel_values),
            (safe_float((brake_cmd or {}).get("brake")), brake_values),
            (safe_float((ctrl_cmd or {}).get("linear_velocity")), target_speed_values),
        ]:
            if value is not None:
                target.append(value)
    _header, _first, last_row = first_and_last_rows(path)
    return {
        "exists": True,
        "rows_sampled": len(rows),
        "sample_truncated": truncated,
        "first_timestamp": row_timestamp_str(rows[0]) if rows else "unknown",
        "last_timestamp": row_timestamp_str(last_row or rows[-1]) if rows else "unknown",
        "steer_cmd_range": min_max(steer_values),
        "accel_cmd_range": min_max(accel_values),
        "brake_cmd_range": min_max(brake_values),
        "target_speed_range": min_max(target_speed_values),
    }


def summarize_object_array_csv(path: Path, object_field: str = "objects") -> dict[str, Any]:
    _fieldnames, rows, truncated = scan_rows(path, limit=CSV_SAMPLE_LIMIT)
    object_counts = []
    labels = set()
    payload_presence = {"roi_image": 0, "pointcloud": 0}
    for row in rows:
        objects = parse_json_cell(row.get(object_field), []) or []
        if not isinstance(objects, list):
            continue
        object_counts.append(len(objects))
        for obj in objects[:10]:
            if not isinstance(obj, dict):
                continue
            if obj.get("label"):
                labels.add(str(obj["label"]))
            pointcloud = obj.get("pointcloud") or {}
            roi_image = obj.get("roi_image") or {}
            if isinstance(pointcloud, dict):
                data = pointcloud.get("data") or {}
                if isinstance(data, dict) and data.get("bytes"):
                    payload_presence["pointcloud"] += 1
            if isinstance(roi_image, dict):
                data = roi_image.get("data") or {}
                if isinstance(data, dict) and data.get("bytes"):
                    payload_presence["roi_image"] += 1
    _header, _first, last_row = first_and_last_rows(path)
    return {
        "exists": True,
        "rows_sampled": len(rows),
        "sample_truncated": truncated,
        "first_timestamp": row_timestamp_str(rows[0]) if rows else "unknown",
        "last_timestamp": row_timestamp_str(last_row or rows[-1]) if rows else "unknown",
        "object_count_range": min_max([float(count) for count in object_counts]),
        "object_count_avg_sample": round(sum(object_counts) / len(object_counts), 3) if object_counts else "unknown",
        "labels_sample": sorted(labels)[:10],
        "embedded_payload_presence_sample": payload_presence,
    }


def summarize_lane_array_csv(path: Path) -> dict[str, Any]:
    _fieldnames, rows, truncated = scan_rows(path, limit=CSV_SAMPLE_LIMIT)
    lane_counts = []
    waypoint_counts = []
    null_rows = 0
    for row in rows:
        lanes = parse_json_cell(row.get("lanes"), []) or []
        if not isinstance(lanes, list):
            null_rows += 1
            continue
        lane_counts.append(len(lanes))
        if lanes:
            first_lane = lanes[0] if isinstance(lanes[0], dict) else {}
            waypoints = (first_lane.get("waypoints") or []) if isinstance(first_lane, dict) else []
            waypoint_counts.append(len(waypoints) if isinstance(waypoints, list) else 0)
        else:
            waypoint_counts.append(0)
    _header, _first, last_row = first_and_last_rows(path)
    return {
        "exists": True,
        "rows_sampled": len(rows),
        "sample_truncated": truncated,
        "first_timestamp": row_timestamp_str(rows[0]) if rows else "unknown",
        "last_timestamp": row_timestamp_str(last_row or rows[-1]) if rows else "unknown",
        "lane_count_range": min_max([float(count) for count in lane_counts]),
        "primary_lane_waypoint_count_range": min_max([float(count) for count in waypoint_counts]),
        "null_rows_sample": null_rows,
    }


def summarize_final_waypoints_csv(path: Path) -> dict[str, Any]:
    _fieldnames, rows, truncated = scan_rows(path, limit=CSV_SAMPLE_LIMIT)
    waypoint_counts = []
    blocked_true = 0
    null_rows = 0
    for row in rows:
        waypoints = parse_json_cell(row.get("waypoints"), []) or []
        if not isinstance(waypoints, list):
            null_rows += 1
            continue
        waypoint_counts.append(len(waypoints))
        if str(row.get("is_blocked", "")).lower() == "true":
            blocked_true += 1
    _header, _first, last_row = first_and_last_rows(path)
    return {
        "exists": True,
        "rows_sampled": len(rows),
        "sample_truncated": truncated,
        "first_timestamp": row_timestamp_str(rows[0]) if rows else "unknown",
        "last_timestamp": row_timestamp_str(last_row or rows[-1]) if rows else "unknown",
        "waypoint_count_range": min_max([float(count) for count in waypoint_counts]),
        "blocked_true_count_sample": blocked_true,
        "null_rows_sample": null_rows,
    }


def summarize_collision_csv(path: Path) -> dict[str, Any]:
    _fieldnames, rows, truncated = scan_rows(path, limit=CSV_SAMPLE_LIMIT)
    other_actor_ids = []
    impulse_magnitudes = []
    for row in rows:
        other_actor = row.get("other_actor_id")
        if other_actor not in (None, ""):
            other_actor_ids.append(str(other_actor))
        impulse = parse_json_cell(row.get("normal_impulse"), {}) or {}
        if isinstance(impulse, dict):
            coords = [safe_float(impulse.get(axis)) for axis in ("x", "y", "z")]
            if all(value is not None for value in coords):
                impulse_magnitudes.append(math.sqrt(sum(value * value for value in coords if value is not None)))
    return {
        "exists": True,
        "rows_sampled": len(rows),
        "sample_truncated": truncated,
        "first_timestamp": row_timestamp_str(rows[0]) if rows else "unknown",
        "last_timestamp": row_timestamp_str(rows[-1]) if rows else "unknown",
        "other_actor_ids_sample": sorted(set(other_actor_ids))[:10],
        "impulse_magnitude_range": min_max(impulse_magnitudes),
    }


def summarize_traffic_lights_csv(path: Path) -> dict[str, Any]:
    _fieldnames, rows, truncated = scan_rows(path, limit=CSV_SAMPLE_LIMIT)
    light_counts = []
    states = set()
    for row in rows:
        lights = parse_json_cell(row.get("traffic_lights"), []) or []
        if not isinstance(lights, list):
            continue
        light_counts.append(len(lights))
        for light in lights[:20]:
            if isinstance(light, dict) and "state" in light:
                states.add(str(light["state"]))
    _header, _first, last_row = first_and_last_rows(path)
    return {
        "exists": True,
        "rows_sampled": len(rows),
        "sample_truncated": truncated,
        "first_timestamp": row_timestamp_str(rows[0]) if rows else "unknown",
        "last_timestamp": row_timestamp_str(last_row or rows[-1]) if rows else "unknown",
        "traffic_light_count_range": min_max([float(count) for count in light_counts]),
        "state_sample": sorted(states),
    }


def summarize_odometry_csv(path: Path) -> dict[str, Any]:
    _fieldnames, rows, truncated = scan_rows(path, limit=CSV_SAMPLE_LIMIT)
    positions = []
    for row in rows:
        pose = parse_json_cell(row.get("pose"), {}) or {}
        if isinstance(pose, dict):
            pose_inner = pose.get("pose") or {}
            position = pose_inner.get("position") or {}
            if isinstance(position, dict):
                positions.append(
                    {
                        "x": safe_float(position.get("x")),
                        "y": safe_float(position.get("y")),
                        "z": safe_float(position.get("z")),
                    }
                )
    _header, _first, last_row = first_and_last_rows(path)
    return {
        "exists": True,
        "rows_sampled": len(rows),
        "sample_truncated": truncated,
        "first_timestamp": row_timestamp_str(rows[0]) if rows else "unknown",
        "last_timestamp": row_timestamp_str(last_row or rows[-1]) if rows else "unknown",
        "start_position": positions[0] if positions else "unknown",
        "end_position_sample": positions[-1] if positions else "unknown",
    }


def summarize_topic(path: Path, kind: str, conversion_root: Path | None) -> dict[str, Any]:
    if kind == "image":
        return summarize_image_csv(path, conversion_root)
    if kind == "pointcloud":
        return summarize_pointcloud_csv(path, conversion_root)
    if kind == "imu":
        return summarize_imu_csv(path)
    if kind == "pose":
        return summarize_pose_csv(path)
    if kind == "vehicle_status":
        return summarize_vehicle_status_csv(path)
    if kind == "vehicle_cmd":
        return summarize_vehicle_cmd_csv(path)
    if kind == "object_array":
        return summarize_object_array_csv(path)
    if kind == "lane_array":
        return summarize_lane_array_csv(path)
    if kind == "final_waypoints":
        return summarize_final_waypoints_csv(path)
    if kind == "collision":
        return summarize_collision_csv(path)
    if kind == "traffic_lights":
        return summarize_traffic_lights_csv(path)
    if kind == "odometry":
        return summarize_odometry_csv(path)
    return {"exists": True, "status": f"no summarizer for {kind}"}


def inspect_convert_script() -> tuple[str, dict[str, Any]]:
    help_run = subprocess.run(
        [sys.executable, "convert_sensor_data_field.py", "-h"],
        cwd=ROOT,
        capture_output=True,
        text=True,
    )
    source = (ROOT / "convert_sensor_data_field.py").read_text(encoding="utf-8")
    report_data = {
        "supports_image_raw": "image_raw" in source,
        "supports_points_raw": "points_raw" in source,
        "supports_object_arrays": "object_array" in source,
        "depends_on": ["numpy", "argparse", "base64", "csv", "json", "pathlib", "sys"],
        "help_ok": help_run.returncode == 0,
        "help_text": help_run.stdout,
    }
    lines = [
        "# convert_sensor_data_field.py Report",
        "",
        "## Help",
        "",
    ]
    if help_run.returncode == 0:
        lines.append("```text")
        lines.append(help_run.stdout.strip())
        lines.append("```")
    else:
        lines.append(f"- `python3 convert_sensor_data_field.py -h` failed with code {help_run.returncode}.")
        lines.append("")
        lines.append("```text")
        lines.append(help_run.stderr.strip())
        lines.append("```")
    lines.extend(
        [
            "",
            "## Usage Summary",
            "",
            f"- Supports `image_raw`: {'yes' if report_data['supports_image_raw'] else 'no'}",
            f"- Supports `points_raw`: {'yes' if report_data['supports_points_raw'] else 'no'}",
            f"- Supports embedded object arrays: {'yes' if report_data['supports_object_arrays'] else 'no'}",
            "- Input path: one extracted topic CSV in single-file mode, or an extracted root directory in batch mode.",
            "- Output path: user-provided `--output` in single-file mode, or `--decoded-root` / per-topic folders in batch mode.",
            "- Original inputs are not modified. The script writes decoded files into output folders, but it may overwrite files in the output folder if rerun with the same filenames.",
            "- Declared dependencies from source inspection: `numpy` plus Python standard library modules (`argparse`, `base64`, `csv`, `json`, `pathlib`, `sys`).",
            "- No direct dependency on `pandas`, `opencv/cv2`, or `rosbag` appears in the script.",
            "",
            "## Notes",
            "",
            "- The script auto-detects `image`, `pointcloud`, and `object_array` CSV layouts from column names.",
            "- For this pipeline, a wrapper script under `scripts/run_conversion_case.py` is used so outputs stay under `processed/converted/` instead of next to the extracted CSV files.",
            "",
        ]
    )
    (PROCESSED / "convert_script_report.md").write_text("\n".join(lines), encoding="utf-8")
    return "\n".join(lines), report_data


def run_case001_conversion() -> dict[str, Any]:
    case_dir = EXTRACTED / "case_001"
    output_root = CONVERTED / "case_001"
    output_root.mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(
        [sys.executable, "scripts/run_conversion_case.py", str(case_dir), str(output_root)],
        cwd=ROOT,
        capture_output=True,
        text=True,
    )
    report = {
        "returncode": completed.returncode,
        "stdout": completed.stdout,
        "stderr": completed.stderr,
        "report_path": output_root / "conversion_report.json",
    }
    error_lines = [
        "# Conversion Errors",
        "",
    ]
    if completed.returncode == 0:
        error_lines.extend(
            [
                "- `case_001` conversion completed successfully.",
                "",
                "```text",
                completed.stdout.strip(),
                "```",
                "",
            ]
        )
    else:
        error_lines.extend(
            [
                f"- `case_001` conversion failed with code {completed.returncode}.",
                "",
                "```text",
                completed.stdout.strip(),
                completed.stderr.strip(),
                "```",
                "",
            ]
        )
    (PROCESSED / "convert_errors.md").write_text("\n".join(error_lines), encoding="utf-8")
    return report


def build_case_summaries(
    cases: list[dict[str, Any]],
    dataset_by_index: dict[int, dict[str, Any]],
    matches_by_case: dict[int, dict[str, TopicMatch]],
) -> dict[int, dict[str, Any]]:
    summaries: dict[int, dict[str, Any]] = {}
    for case in cases:
        case_id = case["case_id"]
        dataset_row = dataset_by_index.get(case_id, {})
        converted_root = CONVERTED / f"case_{case_id:03d}"
        topic_summaries: dict[str, dict[str, Any]] = {}
        available_topics = []
        missing_topics = []
        for topic_name, match in matches_by_case[case_id].items():
            if not match.found:
                missing_topics.append(topic_name)
                continue
            available_topics.append(topic_name)
            rel_path = Path(match.matched_file_path)
            conversion_dir = converted_root / rel_path.stem if rel_path.suffix == ".csv" else None
            topic_summaries[topic_name] = summarize_topic(case["case_dir"] / rel_path, match.kind, conversion_dir)
        summaries[case_id] = {
            "dataset_row": dataset_row,
            "topic_summaries": topic_summaries,
            "available_topics": available_topics,
            "missing_topics": missing_topics,
        }
    return summaries


def event_label_dict(dataset_row: dict[str, Any]) -> dict[str, int]:
    return {
        "collision": int(dataset_row.get("Collision", 0) or 0),
        "stuck": int(dataset_row.get("Stuck", 0) or 0),
        "lane_invasion": int(dataset_row.get("Lane Invasion", 0) or 0),
        "red_light": int(dataset_row.get("Red", 0) or 0),
    }


def failure_timestamp(summary: dict[str, Any], error_data: dict[str, Any]) -> str:
    collision = summary["topic_summaries"].get("/carla/ego_vehicle/collision")
    if collision and collision.get("first_timestamp") not in (None, "unknown"):
        return collision["first_timestamp"]
    if (error_data.get("events") or {}).get("crash"):
        return "unknown (crash event present in error.json, no precise topic timestamp)"
    return "unknown"


def scenario_config_from_error(error_data: dict[str, Any]) -> dict[str, Any]:
    seed = error_data.get("seed") or {}
    weather = error_data.get("weather") or {}
    actors = error_data.get("actors") or []
    puddles = error_data.get("puddles") or []
    mission = "unknown"
    if isinstance(seed, dict):
        if {"sp_x", "sp_y", "wp_x", "wp_y"}.issubset(seed.keys()):
            mission = {
                "start": {key: seed.get(key) for key in ["sp_x", "sp_y", "sp_z", "yaw"]},
                "goal": {key: seed.get(key) for key in ["wp_x", "wp_y", "wp_z", "wp_yaw"]},
            }
    return {
        "map": seed.get("map", "unknown") if isinstance(seed, dict) else "unknown",
        "mission": mission,
        "actors": f"{len(actors)} actor(s)" if isinstance(actors, list) else "unknown",
        "weather": weather if isinstance(weather, dict) else "unknown",
        "puddles": f"{len(puddles)} puddle region(s)" if isinstance(puddles, list) else "unknown",
        "seed": seed if isinstance(seed, dict) else "unknown",
    }


def preliminary_llm_summary(case: dict[str, Any], case_summary: dict[str, Any]) -> str:
    error_data = case["error"]
    labels = event_label_dict(case_summary["dataset_row"])
    topic_summaries = case_summary["topic_summaries"]
    parts = [
        f"Case {case['case_id']} is labeled `{case_summary['dataset_row'].get('Group', 'unknown')}` in Dataset.xlsx.",
        f"Scenario seed map is `{(error_data.get('seed') or {}).get('map', 'unknown')}` with `{len(error_data.get('actors') or [])}` actor(s) and `{len(error_data.get('puddles') or [])}` puddle region(s).",
    ]
    collision = topic_summaries.get("/carla/ego_vehicle/collision")
    if collision:
        parts.append(f"Collision topic exists and first collision timestamp is `{collision.get('first_timestamp', 'unknown')}`.")
    else:
        parts.append("No explicit collision topic was found for this case.")
    image = topic_summaries.get("/image_raw")
    points = topic_summaries.get("/points_raw")
    if image:
        parts.append(
            f"image_raw exists with shape {image.get('height', 'unknown')}x{image.get('width', 'unknown')} and decode status `{image.get('decode_status', 'unknown')}`."
        )
    if points:
        parts.append(
            f"points_raw exists with width {points.get('width', 'unknown')} and fields {points.get('fields', [])[:6]}."
        )
    detection = topic_summaries.get("/detection/fusion_tools/objects")
    if detection:
        parts.append(
            f"Detection objects are available with sampled object count range `{detection.get('object_count_range', 'unknown')}` and labels {detection.get('labels_sample', [])}."
        )
    planning = topic_summaries.get("/final_waypoints")
    if planning:
        parts.append(
            f"final_waypoints are available with sampled waypoint count range `{planning.get('waypoint_count_range', 'unknown')}`."
        )
    actuation = topic_summaries.get("/carla/ego_vehicle/vehicle_status")
    if actuation:
        parts.append(
            f"Vehicle status speed range sample is `{actuation.get('velocity_range', 'unknown')}`."
        )
    red_event = (error_data.get("events") or {}).get("red")
    lane_event = (error_data.get("events") or {}).get("lane_invasion")
    stuck_event = (error_data.get("events") or {}).get("stuck")
    parts.append(
        f"error.json events report red={red_event}, lane_invasion={lane_event}, stuck={stuck_event}, while dataset labels are collision={labels['collision']}, stuck={labels['stuck']}, lane_invasion={labels['lane_invasion']}, red_light={labels['red_light']}."
    )
    return " ".join(parts)


def case_card_markdown(case: dict[str, Any], summary: dict[str, Any]) -> str:
    case_id = case["case_id"]
    dataset_row = summary["dataset_row"]
    labels = event_label_dict(dataset_row)
    error_data = case["error"]
    scenario = scenario_config_from_error(error_data)
    topic_summaries = summary["topic_summaries"]
    collision_summary = topic_summaries.get("/carla/ego_vehicle/collision")
    image_summary = topic_summaries.get("/image_raw", {})
    points_summary = topic_summaries.get("/points_raw", {})
    detection_summary = topic_summaries.get("/detection/fusion_tools/objects", {})
    prediction_summary = topic_summaries.get("/prediction/motion_predictor/objects", {})
    lane_summary = topic_summaries.get("/lane_waypoints_array", {})
    final_summary = topic_summaries.get("/final_waypoints", {})
    vehicle_cmd_summary = topic_summaries.get("/vehicle_cmd", {})
    vehicle_status_summary = topic_summaries.get("/carla/ego_vehicle/vehicle_status", {})
    odometry_summary = topic_summaries.get("/carla/ego_vehicle/odometry", {})
    traffic_summary = topic_summaries.get("/carla/traffic_lights", {})
    objects_summary = topic_summaries.get("/carla/objects", {})
    imu_summary = topic_summaries.get("/carla/ego_vehicle/imu/imu", {})
    current_pose_summary = topic_summaries.get("/current_pose", {})
    missing_topics = summary["missing_topics"]

    lines = [
        f"# case_{case_id:03d}",
        "",
        "## 1. Basic info",
        "",
        f"- case_id: `case_{case_id:03d}`",
        f"- zip file: `{case_id}.zip`",
        f"- fault group from Dataset.xlsx: `{dataset_row.get('Group', 'unknown')}` ({FAULT_GROUP_MEANINGS.get(dataset_row.get('Group', ''), 'unknown meaning')})",
        f"- Collision / Stuck / Lane Invasion / Red labels: `{labels['collision']}` / `{labels['stuck']}` / `{labels['lane_invasion']}` / `{labels['red_light']}`",
        f"- available topics: {', '.join(summary['available_topics']) if summary['available_topics'] else 'none'}",
        "",
        "## 2. Scenario-level data",
        "",
        f"- map / town: `{scenario['map']}`",
        f"- mission: `{short_json(scenario['mission'], 220) if scenario['mission'] != 'unknown' else 'unknown'}`",
        f"- seed info: `{short_json(scenario['seed'], 260) if scenario['seed'] != 'unknown' else 'unknown'}`",
        f"- actors: `{short_json(error_data.get('actors'), 220)}`",
        f"- weather: `{short_json(error_data.get('weather'), 220)}`",
        f"- puddles: `{short_json(error_data.get('puddles'), 220)}`",
        "",
        "## 3. Oracle evidence",
        "",
        f"- collision exists: `{'yes' if collision_summary else 'no'}`",
        f"- red-light violation exists: `dataset={labels['red_light']}`, `error.json={(error_data.get('events') or {}).get('red', 'unknown')}`",
        f"- lane invasion exists: `dataset={labels['lane_invasion']}`, `error.json={(error_data.get('events') or {}).get('lane_invasion', 'unknown')}`",
        f"- stuck / immobility exists: `dataset={labels['stuck']}`, `error.json={(error_data.get('events') or {}).get('stuck', 'unknown')}`",
        f"- failure timestamp: `{failure_timestamp(summary, error_data)}`",
        f"- collision detail sample: `{short_json(collision_summary, 220) if collision_summary else 'unknown'}`",
        "",
        "## 4. Sensing evidence",
        "",
        f"- image_raw: `{short_json(image_summary, 260) if image_summary else 'not found'}`",
        f"- points_raw: `{short_json(points_summary, 260) if points_summary else 'not found'}`",
        f"- IMU: `{short_json(imu_summary, 220) if imu_summary else 'not found'}`",
        "",
        "## 5. Perception evidence",
        "",
        f"- detection/fusion objects: `{short_json(detection_summary, 260) if detection_summary else 'not found'}`",
        f"- current_pose: `{short_json(current_pose_summary, 220) if current_pose_summary else 'not found'}`",
        f"- prediction objects: `{short_json(prediction_summary, 260) if prediction_summary else 'not found'}`",
        "",
        "## 6. Planning evidence",
        "",
        f"- lane_waypoints_array: `{short_json(lane_summary, 240) if lane_summary else 'not found'}`",
        f"- final_waypoints: `{short_json(final_summary, 240) if final_summary else 'not found'}`",
        "",
        "## 7. Actuation evidence",
        "",
        f"- vehicle_cmd: `{short_json(vehicle_cmd_summary, 240) if vehicle_cmd_summary else 'not found'}`",
        f"- vehicle_status: `{short_json(vehicle_status_summary, 240) if vehicle_status_summary else 'not found'}`",
        "",
        "## 8. Environment evidence",
        "",
        f"- traffic_lights: `{short_json(traffic_summary, 220) if traffic_summary else 'not found'}`",
        f"- objects: `{short_json(objects_summary, 220) if objects_summary else 'not found'}`",
        f"- odometry: `{short_json(odometry_summary, 220) if odometry_summary else 'not found'}`",
        "",
        "## 9. Preliminary causal explanation input",
        "",
        preliminary_llm_summary(case, summary),
        "",
        "## 10. Missing data / decoding problems",
        "",
        f"- Missing topics: {', '.join(missing_topics) if missing_topics else 'none'}",
        f"- image_raw decode status: `{image_summary.get('decode_status', 'unknown') if image_summary else 'not found'}`",
        f"- points_raw decode status: `{points_summary.get('decode_status', 'unknown') if points_summary else 'not found'}`",
        "- detection / prediction object arrays may still contain embedded payloads that need dedicated decoding if a later analysis requires per-object ROI images or point clouds.",
        "- Dataset labels and `error.json` events are not fully aligned for some cases; manual confirmation is recommended before treating non-collision labels as oracle ground truth.",
        "",
    ]
    return "\n".join(lines)


def llm_input_json(case: dict[str, Any], summary: dict[str, Any]) -> dict[str, Any]:
    dataset_row = summary["dataset_row"]
    labels = event_label_dict(dataset_row)
    scenario = scenario_config_from_error(case["error"])
    topic_summaries = summary["topic_summaries"]
    return {
        "case_id": f"case_{case['case_id']:03d}",
        "fault_label": {
            **labels,
            "group": dataset_row.get("Group", "unknown"),
        },
        "scenario_config": {
            "map": scenario["map"],
            "mission": scenario["mission"],
            "actors": scenario["actors"],
            "weather": scenario["weather"],
            "puddles": scenario["puddles"],
        },
        "available_topics": summary["available_topics"],
        "trace_summary": {
            "sensing": {
                "image_raw": topic_summaries.get("/image_raw"),
                "points_raw": topic_summaries.get("/points_raw"),
                "imu": topic_summaries.get("/carla/ego_vehicle/imu/imu"),
            },
            "perception": {
                "detection_fusion_objects": topic_summaries.get("/detection/fusion_tools/objects"),
                "current_pose": topic_summaries.get("/current_pose"),
                "prediction_objects": topic_summaries.get("/prediction/motion_predictor/objects"),
            },
            "planning": {
                "lane_waypoints_array": topic_summaries.get("/lane_waypoints_array"),
                "final_waypoints": topic_summaries.get("/final_waypoints"),
            },
            "actuation": {
                "vehicle_cmd": topic_summaries.get("/vehicle_cmd"),
                "vehicle_status": topic_summaries.get("/carla/ego_vehicle/vehicle_status"),
            },
            "oracle": {
                "collision": topic_summaries.get("/carla/ego_vehicle/collision"),
                "traffic_lights": topic_summaries.get("/carla/traffic_lights"),
                "odometry": topic_summaries.get("/carla/ego_vehicle/odometry"),
                "objects": topic_summaries.get("/carla/objects"),
                "error_json_events": case["error"].get("events"),
            },
        },
        "known_limitations": [
            "Large CSV topics were summarized with bounded scans for reproducibility and memory safety.",
            "Only case_001 was fully decoded with convert_sensor_data_field.py in this pipeline run.",
            "Dataset labels and error.json oracle events are not fully aligned for some cases.",
        ],
    }


def write_case_outputs(cases: list[dict[str, Any]], summaries: dict[int, dict[str, Any]]) -> None:
    for case in cases:
        case_id = case["case_id"]
        summary = summaries[case_id]
        card_path = CASE_CARDS / f"case_{case_id:03d}.md"
        card_path.write_text(case_card_markdown(case, summary), encoding="utf-8")
        json_path = LLM_INPUTS / f"case_{case_id:03d}.json"
        json_path.write_text(json.dumps(llm_input_json(case, summary), indent=2, ensure_ascii=False), encoding="utf-8")


def write_prompt_template() -> None:
    content = """# Causal Explanation Prompt Template

You are given one DriveFuzz / Autoware failure case in structured form.

Use only the evidence provided in:

- `scenario summary`
- `fault label`
- `available topics`
- `trace summary`
- `known limitations`

Do not invent missing facts. If the evidence is insufficient, write `unknown`.

## Input

```json
{{CASE_JSON_HERE}}
```

## Task

Produce the following fields:

1. `fault_layer`
Allowed values: `sensing`, `perception`, `planning`, `actuation`, `simulator`, `map`, `unknown`

2. `causal_chain`
Write a short step-by-step explanation from observable evidence to failure outcome.

3. `supporting_evidence`
List only evidence directly supported by the input data.

4. `uncertain_evidence`
List gaps, contradictions, or weak signals.

5. `possible_root_cause_pattern`
Summarize the most plausible reusable failure pattern, or `unknown`.

6. `bug2scenario_suitability`
State whether this case is suitable for Bug2Scenario-style root-cause-preserving scenario amplification.
Allowed values: `yes`, `no`, `uncertain`
Explain why.

7. `possible_validation_rules`
List validation rules that could be checked in replay or regenerated scenarios.

## Constraints

- Do not use background knowledge that is not present in the input.
- Do not assume causal direction without evidence.
- If dataset labels and error traces disagree, call that out explicitly.
- Prefer concise technical language.
"""
    (PROMPTS / "causal_explanation_prompt_template.md").write_text(content, encoding="utf-8")


def write_workflows_and_metrics() -> None:
    content = """# Workflows And Metrics

## A. Proposal workflow: Bug2Scenario / root-cause-preserving scenario amplification

- Input: case card + fault label + selected topic summaries + root-cause evidence
- Output: causal explanation, root-cause pattern, scenario templates, validation rules
- Metrics:
- same-root-cause reproduction rate
- semantic consistency rate
- executability rate
- physical validity rate
- LLM hallucination rate
- query number
- simulator execution saved

## B. Pure-LLM workflow

- Input: scenario data + fault label
- Output: LLM causal explanation and generated scenarios
- Metrics:
- layer accuracy
- causal chain quality
- hallucination rate
- scenario validity

## C. Previous work / human annotation workflow

- Input: raw ROS topics or case card
- Output: human causal explanation / root-cause label
- Metrics:
- annotation time
- agreement with known fault label
- evidence completeness
"""
    (PROCESSED / "workflows_and_metrics.md").write_text(content, encoding="utf-8")


def write_readme_report(cases: list[dict[str, Any]], summaries: dict[int, dict[str, Any]]) -> None:
    lines = [
        "# Pipeline Report",
        "",
        "## Processed Cases",
        "",
    ]
    for case in cases:
        case_id = case["case_id"]
        summary = summaries[case_id]
        dataset_row = summary["dataset_row"]
        image = summary["topic_summaries"].get("/image_raw", {})
        points = summary["topic_summaries"].get("/points_raw", {})
        detection = summary["topic_summaries"].get("/detection/fusion_tools/objects")
        prediction = summary["topic_summaries"].get("/prediction/motion_predictor/objects")
        lines.extend(
            [
                f"### case_{case_id:03d}",
                "",
                f"- fault group: `{dataset_row.get('Group', 'unknown')}`",
                f"- topics found: {', '.join(summary['available_topics']) if summary['available_topics'] else 'none'}",
                f"- topics not found: {', '.join(summary['missing_topics']) if summary['missing_topics'] else 'none'}",
                f"- image_raw decode: `{image.get('decode_status', 'unknown') if image else 'not found'}`",
                f"- points_raw decode: `{points.get('decode_status', 'unknown') if points else 'not found'}`",
                f"- detection needs extra decoding: `{'yes' if detection and detection.get('embedded_payload_presence_sample', {}).get('roi_image', 0) else 'possible / not required for current summary'}`",
                f"- prediction needs extra decoding: `{'yes' if prediction and prediction.get('embedded_payload_presence_sample', {}).get('pointcloud', 0) else 'possible / not required for current summary'}`",
                "",
            ]
        )
    lines.extend(
        [
            "## Recommendation",
            "",
            "- Next case to process: `13.zip`",
            "- Reason: it is smaller than `24.zip`, `27.zip`, and `37.zip`, follows the same structure, and includes an explicit collision topic plus a composite `C+L` label.",
            "",
            "## Is Current Data Enough For LLM Causal Explanation?",
            "",
            "- `case_001`: mostly yes for a first-pass LLM explanation because oracle, planning, actuation, perception, and decoded sensing outputs are available.",
            "- `case_003` and `case_007`: partially yes. There is enough structured context for a tentative explanation, but red-light / lane-invasion evidence needs manual confirmation because dataset labels do not fully match `error.json` events and sensing topics were not fully decoded in this run.",
            "",
        ]
    )
    (PROCESSED / "README_pipeline_report.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    ensure_dirs()
    write_directory_and_zip_summary()
    dataset_rows, dataset_by_index, dataset_markdown = load_dataset()
    write_dataset_summary(dataset_rows, dataset_markdown)
    extract_selected_cases()
    cases = [load_case_static(case_id) for case_id in TARGET_CASES]
    write_file_inventory(cases)
    matches_by_case = write_topic_availability(cases)
    inspect_convert_script()
    run_case001_conversion()
    summaries = build_case_summaries(cases, dataset_by_index, matches_by_case)
    write_case_outputs(cases, summaries)
    write_prompt_template()
    write_workflows_and_metrics()
    write_readme_report(cases, summaries)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
